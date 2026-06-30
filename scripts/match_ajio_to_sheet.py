import re
import openpyxl

XLSX_PATH = "data/sku_mapping.xlsx"
MATCH_THRESHOLD = 0.4

# Paste scraped Ajio products here as (name, price, url) tuples.
AJIO_PRODUCTS = [
    ("Ten X You Men All-Rounder Cricket Shoes", "₹3,279", "https://www.ajio.com/ten-x-you-men-all-rounder-cricket-shoes/p/469817211_white?"),
    ("Ten X You Men All-Rounder Cricket Shoes", "₹3,279", "https://www.ajio.com/ten-x-you-men-all-rounder-cricket-shoes/p/469817211_limegreen?"),
    ("Ten X You Unisex Logo Print Reset Slides", "₹1,499", "https://www.ajio.com/ten-x-you-unisex-logo-print-reset-slides/p/469817214_black?"),
    ("Ten X You Unisex Logo Print Reset Slides", "₹1,094", "https://www.ajio.com/ten-x-you-unisex-logo-print-reset-slides/p/469817214_limegreen?"),
    ("Ten X You Unisex Crossover Slip-On Sneakers", "₹3,420", "https://www.ajio.com/ten-x-you-unisex-crossover-slip-on-sneakers/p/469817208_green?"),
    ("Ten X You Unisex Switch Fan Edit 2.0 Multi-Sport Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-switch-fan-edit-2-0-multi-sport-sneakers/p/469817213_blue?"),
    ("Ten X You Unisex Crossover Slip-On Sneakers", "₹3,420", "https://www.ajio.com/ten-x-you-unisex-crossover-slip-on-sneakers/p/469817208_lilac?"),
    ("Ten X You Unisex Switch Fan Edit 2.0 Multi-Sport Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-switch-fan-edit-2-0-multi-sport-sneakers/p/469817213_yellow?"),
    ("Ten X You Unisex Sundowner Denim Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-sundowner-denim-sneakers/p/469817209_green?"),
    ("Ten X You Unisex Switch Fan Edit 2.0 Multi-Sport Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-switch-fan-edit-2-0-multi-sport-sneakers/p/469817213_orange?"),
    ("Ten X You Unisex Aeonic Running Shoes", "₹3,960", "https://www.ajio.com/ten-x-you-unisex-aeonic-running-shoes/p/469817206_green?"),
    ("Ten X You Unisex Crossover Slip-On Sneakers", "₹3,420", "https://www.ajio.com/ten-x-you-unisex-crossover-slip-on-sneakers/p/469817208_brown?"),
    ("Ten X You Men NOX Legacy Sneakers", "₹2,940", "https://www.ajio.com/ten-x-you-men-nox-legacy-sneakers/p/469817210_black?"),
    ("Ten X You Unisex Switch Fan Edit 2.0 Multi-Sport Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-switch-fan-edit-2-0-multi-sport-sneakers/p/469817213_purple?"),
    ("Ten X You Unisex Switch Fan Edit 2.0 Multi-Sport Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-switch-fan-edit-2-0-multi-sport-sneakers/p/469817213_red?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹680", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817539_blue?"),
    ("Ten X You Men Regular Fit Crew-Neck Running Singlet", "₹880", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-singlet/p/469817541_blue?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817536_grey?"),
    ("Ten X You Men Centurion Pro Cricket Shoes", "₹8,969", "https://www.ajio.com/ten-x-you-men-centurion-pro-cricket-shoes/p/469817203_white?"),
    ("Ten X You Unisex Sundowner Denim Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-sundowner-denim-sneakers/p/469817209_black?"),
    ("Ten X You Women Regular Fit Crew-Neck Running Tank T-Shirt", "₹480", "https://www.ajio.com/ten-x-you-women-regular-fit-crew-neck-running-tank-t-shirt/p/469817556_blue?"),
    ("Ten X You Women Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-women-relaxed-fit-running-shorts/p/469817551_black?"),
    ("Ten X You Unisex Aeonic Running Shoes", "₹3,960", "https://www.ajio.com/ten-x-you-unisex-aeonic-running-shoes/p/469817206_navy?"),
    ("Ten X You Unisex Zenflo Walking Shoes", "₹4,959", "https://www.ajio.com/ten-x-you-unisex-zenflo-walking-shoes/p/469817207_black?"),
    ("Ten X You Unisex Crossover Slip-On Sneakers", "₹4,949", "https://www.ajio.com/ten-x-you-unisex-crossover-slip-on-sneakers/p/469817208_blackwhite?"),
    ("Ten X You Unisex Crossover Slip-On Sneakers", "₹5,039", "https://www.ajio.com/ten-x-you-unisex-crossover-slip-on-sneakers/p/469817208_black?"),
    ("Ten X You Unisex Aeonic Running Shoes", "₹3,960", "https://www.ajio.com/ten-x-you-unisex-aeonic-running-shoes/p/469817206_red?"),
    ("Ten X You Unisex Aeonic Running Shoes", "₹3,960", "https://www.ajio.com/ten-x-you-unisex-aeonic-running-shoes/p/469817206_black?"),
    ("Ten X You Men Centurion Pro Cricket Shoes", "₹8,969", "https://www.ajio.com/ten-x-you-men-centurion-pro-cricket-shoes/p/469817203_blue?"),
    ("Ten X You Unisex Zenflo Walking Shoes", "₹4,959", "https://www.ajio.com/ten-x-you-unisex-zenflo-walking-shoes/p/469817207_red?"),
    ("Ten X You Unisex Sundowner Denim Sneakers", "₹2,450", "https://www.ajio.com/ten-x-you-unisex-sundowner-denim-sneakers/p/469817209_blue?"),
    ("Ten X You Men Regular Fit Shorts", "₹1,170", "https://www.ajio.com/ten-x-you-men-regular-fit-shorts/p/469817548_black?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹680", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817539_black?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹2,299", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817536_orange?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹680", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817539_red?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹680", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817539_green?"),
    ("Ten X You Men Regular Fit Crew-Neck Singlet", "₹680", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-singlet/p/469817543_red?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817536_black?"),
    ("Ten X You Men Regular Fit Running Shorts", "₹1,290", "https://www.ajio.com/ten-x-you-men-regular-fit-running-shorts/p/469817547_greymelange?"),
    ("Ten X You Men Typographic Print Oversized Fit Crew-Neck T-Shirt", "₹1,080", "https://www.ajio.com/ten-x-you-men-typographic-print-oversized-fit-crew-neck-t-shirt/p/469817540_mauve?"),
    ("Ten X You Men Regular Fit Crew-Neck Running Singlet", "₹690", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-singlet/p/469817542_blue?"),
    ("Ten X You Men Regular Fit Crew-Neck Training T-Shirt", "₹897", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-training-t-shirt/p/469817537_blue?"),
    ("Ten X You Men Regular Fit Crew-Neck Running T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-t-shirt/p/469817536_red?"),
    ("Ten X You Men Typographic Print Oversized Fit Crew-Neck T-Shirt", "₹1,110", "https://www.ajio.com/ten-x-you-men-typographic-print-oversized-fit-crew-neck-t-shirt/p/469817540_offwhite?"),
    ("Ten X You Men Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-men-relaxed-fit-running-shorts/p/469817546_black?"),
    ("Ten X You Men Regular Fit Crew-Neck Training T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-training-t-shirt/p/469817535_green?"),
    ("Ten X You Men Regular Fit Crew-Neck Running Singlet", "₹880", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-running-singlet/p/469817541_black?"),
    ("Ten X You Men Regular Fit Running Shorts", "₹1,290", "https://www.ajio.com/ten-x-you-men-regular-fit-running-shorts/p/469817547_grey?"),
    ("Ten X You Women Regular Fit Crew-Neck Running T-Shirt", "₹2,299", "https://www.ajio.com/ten-x-you-women-regular-fit-crew-neck-running-t-shirt/p/469817555_orange?"),
    ("Ten X You Women Logo Print Slim Fit Crew-Neck T-Shirt", "₹792", "https://www.ajio.com/ten-x-you-women-logo-print-slim-fit-crew-neck-t-shirt/p/469817560_brown?"),
    ("Ten X You Women Typographic Print Regular Fit Boat-Neck T-Shirt", "₹704", "https://www.ajio.com/ten-x-you-women-typographic-print-regular-fit-boat-neck-t-shirt/p/469817559_blue?"),
    ("Ten X You Women Regular Fit Crew-Neck Running T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-women-regular-fit-crew-neck-running-t-shirt/p/469817555_mauve?"),
    ("Ten X You Women Crew-Neck Regular Fit T-Shirt", "₹1,999", "https://www.ajio.com/ten-x-you-women-crew-neck-regular-fit-t-shirt/p/469817561_orange?"),
    ("Ten X You Women Ribbed Regular Fit Tank Top", "₹700", "https://www.ajio.com/ten-x-you-women-ribbed-regular-fit-tank-top/p/469817562_copper?"),
    ("Ten X You Women Typographic Print Regular Fit Boat-Neck T-Shirt", "₹682", "https://www.ajio.com/ten-x-you-women-typographic-print-regular-fit-boat-neck-t-shirt/p/469817559_seagreen?"),
    ("Ten X You Women Typographic Print Oversized Fit Crew-Neck T-Shirt", "₹1,110", "https://www.ajio.com/ten-x-you-women-typographic-print-oversized-fit-crew-neck-t-shirt/p/469817558_powderpink?"),
    ("Ten X You Women Regular Fit Crew-Neck Running T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-women-regular-fit-crew-neck-running-t-shirt/p/469817555_black?"),
    ("Ten X You Women Regular Fit Round-Neck Running T-Shirt", "₹483", "https://www.ajio.com/ten-x-you-women-regular-fit-round-neck-running-t-shirt/p/469817557_black?"),
    ("Ten X You Women Regular Fit Round-Neck Training T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-women-regular-fit-round-neck-training-t-shirt/p/469817554_mauve?"),
    ("Ten X You Women Regular Fit Round-Neck Running T-Shirt", "₹483", "https://www.ajio.com/ten-x-you-women-regular-fit-round-neck-running-t-shirt/p/469817557_blue?"),
    ("Ten X You Women Typographic Print Regular Fit Boat-Neck T-Shirt", "₹704", "https://www.ajio.com/ten-x-you-women-typographic-print-regular-fit-boat-neck-t-shirt/p/469817559_violet?"),
    ("Ten X You Women Logo Print Slim Fit Crew-Neck T-Shirt", "₹792", "https://www.ajio.com/ten-x-you-women-logo-print-slim-fit-crew-neck-t-shirt/p/469817560_purple?"),
    ("Ten X You Women Regular Fit Crew-Neck Running Tank T-Shirt", "₹480", "https://www.ajio.com/ten-x-you-women-regular-fit-crew-neck-running-tank-t-shirt/p/469817556_black?"),
    ("Ten X You Women Ribbed Regular Fit Tank Top", "₹700", "https://www.ajio.com/ten-x-you-women-ribbed-regular-fit-tank-top/p/469817562_black?"),
    ("Ten X You Men Regular Fit Crew-Neck Training T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-training-t-shirt/p/469817535_purple?"),
    ("Ten X You Men Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-men-relaxed-fit-running-shorts/p/469817545_black?"),
    ("Ten X You Men Typographic Print Regular Fit Crew-Neck Singlet", "₹680", "https://www.ajio.com/ten-x-you-men-typographic-print-regular-fit-crew-neck-singlet/p/469817544_black?"),
    ("Ten X You Women Training Shorts", "₹1,140", "https://www.ajio.com/ten-x-you-women-training-shorts/p/469817552_blue?"),
    ("Ten X You Women Paneled Relaxed Fit Walking Shorts", "₹1,200", "https://www.ajio.com/ten-x-you-women-paneled-relaxed-fit-walking-shorts/p/469817549_brown?"),
    ("Ten X You Women Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-women-relaxed-fit-running-shorts/p/469817550_black?"),
    ("Ten X You Women Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-women-relaxed-fit-running-shorts/p/469817550_red?"),
    ("Ten X You Women Ribbed Regular Fit Tank Top", "₹680", "https://www.ajio.com/ten-x-you-women-ribbed-regular-fit-tank-top/p/469817562_pink?"),
    ("Ten X You Women Regular Fit Round-Neck Training T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-women-regular-fit-round-neck-training-t-shirt/p/469817554_grey?"),
    ("Ten X You Women Regular Fit Round-Neck Training T-Shirt", "₹989", "https://www.ajio.com/ten-x-you-women-regular-fit-round-neck-training-t-shirt/p/469817554_pink?"),
    ("Ten X You Men Relaxed Fit Running Shorts", "₹1,470", "https://www.ajio.com/ten-x-you-men-relaxed-fit-running-shorts/p/469817545_blue?"),
    ("Ten X You Men Regular Fit Crew-Neck Training T-Shirt", "₹2,299", "https://www.ajio.com/ten-x-you-men-regular-fit-crew-neck-training-t-shirt/p/469817535_blue?"),
    ("Ten X You Women Typographic Print Regular Fit Boat-Neck T-Shirt", "₹704", "https://www.ajio.com/ten-x-you-women-typographic-print-regular-fit-boat-neck-t-shirt/p/469817559_black?"),
    ("Ten X You Women Typographic Print Oversized Fit Crew-Neck T-Shirt", "₹2,999", "https://www.ajio.com/ten-x-you-women-typographic-print-oversized-fit-crew-neck-t-shirt/p/469817558_pink?"),
    ("Ten X You Women Paneled Relaxed Fit Walking Shorts", "₹1,170", "https://www.ajio.com/ten-x-you-women-paneled-relaxed-fit-walking-shorts/p/469817549_black?"),
]


def normalize(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return set(text.split())


def overlap_score(words_a, words_b):
    if not words_a or not words_b:
        return 0.0
    intersection = words_a & words_b
    union = words_a | words_b
    return len(intersection) / len(union)


def find_best_match(product_name, ajio_products):
    target_words = normalize(product_name)

    best_score = 0.0
    best_match = None

    for name, price, url in ajio_products:
        score = overlap_score(target_words, normalize(name))
        if score > best_score:
            best_score = score
            best_match = (name, price, url)

    if best_match and best_score >= MATCH_THRESHOLD:
        return best_match, best_score
    return None, best_score


def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers)
    sku_col = headers.index("SKU")
    name_col = headers.index("Product Name")

    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[sku_col]
        product_name = row[name_col]
        if not sku or not product_name:
            continue

        match, score = find_best_match(product_name, AJIO_PRODUCTS)

        if match:
            name, price, url = match
            print(f"{sku} | {product_name}")
            print(f"  -> MATCH ({score:.2f}): {name} | {price} | {url}")
        else:
            print(f"{sku} | {product_name}")
            print(f"  -> NO MATCH (best score: {score:.2f})")
        print()


if __name__ == "__main__":
    main()
