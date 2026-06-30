import json
import re
from collections import Counter

import openpyxl

XLSX_PATH = "data/sku_mapping.xlsx"
MYNTRA_RESULTS_PATH = "data/myntra_search_results.json"
AJIO_RESULTS_PATH = "data/ajio_search_results.json"
MATCH_THRESHOLD = 0.4


def normalize(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return set(text.split())


def overlap_score(words_a, words_b):
    if not words_a or not words_b:
        return 0.0
    intersection = words_a & words_b
    union = words_a | words_b
    return len(intersection) / len(union)


def find_best_match(product_name, candidates):
    target_words = normalize(product_name)

    scored = [(overlap_score(target_words, normalize(c["name"])), c) for c in candidates]
    if not scored:
        return None, 0.0

    best_score = max(score for score, _ in scored)
    if best_score < MATCH_THRESHOLD:
        return None, best_score

    tied = [c for score, c in scored if score == best_score]
    if len(tied) == 1:
        return tied[0], best_score

    price_counts = Counter(c["price"] for c in tied)
    max_count = max(price_counts.values())
    modes = [price for price, count in price_counts.items() if count == max_count]

    if max_count > 1 and len(modes) == 1:
        mode_price = modes[0]
        for c in tied:
            if c["price"] == mode_price:
                return c, best_score

    return tied[0], best_score


def truncate(text, length):
    text = text or "—"
    return text if len(text) <= length else text[: length - 1] + "…"


def main():
    with open(MYNTRA_RESULTS_PATH, encoding="utf-8") as f:
        myntra_products = json.load(f)
    with open(AJIO_RESULTS_PATH, encoding="utf-8") as f:
        ajio_products = json.load(f)

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    sku_col = headers.index("SKU")
    name_col = headers.index("Product Name")

    col_widths = (14, 32, 32, 7, 32, 7)
    header_row = ("SKU", "Product Name", "Matched Myntra Name", "M-Score", "Matched Ajio Name", "A-Score")
    print(" | ".join(h.ljust(w) for h, w in zip(header_row, col_widths)))
    print("-+-".join("-" * w for w in col_widths))

    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[sku_col]
        product_name = row[name_col]
        if not sku or not product_name:
            continue

        myntra_match, myntra_score = find_best_match(product_name, myntra_products)
        ajio_match, ajio_score = find_best_match(product_name, ajio_products)

        myntra_name = myntra_match["name"] if myntra_match else None
        ajio_name = ajio_match["name"] if ajio_match else None

        cells = (
            truncate(str(sku), col_widths[0]),
            truncate(product_name, col_widths[1]),
            truncate(myntra_name, col_widths[2]),
            f"{myntra_score:.2f}",
            truncate(ajio_name, col_widths[4]),
            f"{ajio_score:.2f}",
        )
        print(" | ".join(c.ljust(w) for c, w in zip(cells, col_widths)))


if __name__ == "__main__":
    main()
