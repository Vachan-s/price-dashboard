import asyncio
import csv
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timezone
from io import StringIO

import openpyxl
from flask import Flask, jsonify, make_response, render_template
from playwright.async_api import async_playwright

from scrapers.amazon_search import scrape_amazon_price_with_page
from scrapers.myntra import scrape_myntra_price_with_page
from scrapers.tenxyou import scrape_tenxyou_price_with_page
from scrapers.tenxyou_search import _product_key as _tenxyou_base_slug

os.chdir(os.path.dirname(os.path.abspath(__file__)))

if sys.platform == "win32":
    subprocess.call(
        'for /f "tokens=5" %a in (\'netstat -ano ^| findstr :5000 ^| findstr LISTENING\') do taskkill /F /PID %a',
        shell=True,
    )

app = Flask(__name__)

XLSX_PATH        = "data/sku_mapping.xlsx"
URL_MAPPING_PATH = "data/url_mapping.xlsx"
RESULTS_PATH     = "data/results.json"

AMAZON_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

FUZZY_MATCH_THRESHOLD = 0.6
SCRAPE_CONCURRENCY    = 5

# ── Price helpers ────────────────────────────────────────────────────────────


def _parse_price(price_str) -> float | None:
    """Extract numeric value from a price string. Returns None if 0 or unparseable."""
    if isinstance(price_str, (int, float)):
        return float(price_str) if price_str else None
    # Match a proper number (with optional thousands commas and decimal fraction)
    # rather than stripping stray punctuation, since labels like "Rs." would
    # otherwise leave a leading decimal point behind (e.g. "Rs. 1199" → ".1199").
    m = re.search(r"\d[\d,]*\.\d+|\d[\d,]*", str(price_str or ""))
    if not m:
        return None
    try:
        v = float(m.group(0).replace(",", ""))
        return v if v else None
    except ValueError:
        return None


def _normalize_price(price_str) -> int | None:
    """Convert any price representation to a plain integer (INR). Returns None if 0 or unparseable."""
    v = _parse_price(price_str)
    return int(v) if v else None


def _norm_url(url: str) -> str:
    """Normalize a URL for equality comparison: strip query string and trailing slash."""
    return (url or "").split("?")[0].rstrip("/")


_AMAZON_ASIN_RE = re.compile(r"/dp/([A-Z0-9]{10})")
_MYNTRA_ID_RE   = re.compile(r"/(\d{6,8})/")
_AJIO_CODE_RE   = re.compile(r"/p/(\d+)_")


def _extract_stable_id(url: str, platform: str) -> str:
    """Extract a stable per-platform product identifier from a URL, so matching
    a mapping-sheet URL against a scraped/cached listing is robust to query
    params, /ref= suffixes, color-variant suffixes, and slug wording
    differences. Falls back to the plain normalized URL if the expected
    pattern isn't found (e.g. a malformed or unrecognized URL)."""
    url = url or ""
    if platform == "amazon":
        m = _AMAZON_ASIN_RE.search(url)
        if m:
            return m.group(1)
    elif platform == "myntra":
        m = _MYNTRA_ID_RE.search(url)
        if m:
            return m.group(1)
    elif platform == "ajio":
        m = _AJIO_CODE_RE.search(url)
        if m:
            return m.group(1)
    elif platform == "tenxyou":
        return _tenxyou_base_slug(url)
    return _norm_url(url)


# ── Fuzzy-matching helpers (supplementary pass for SKUs still missing a Myntra
# or TenXYou price after primary per-URL scraping — see _run_scrape) ─────────


_FUZZY_STOP_WORDS = {"ten", "x", "you", "unisex", "women", "men", "fit", "regular"}


def _stem(word: str) -> str:
    return word[:-1] if word.endswith("s") and len(word) > 3 else word


def normalize(text: str) -> set:
    """Normalize a product name into a stemmed word set, stripping brand/generic
    noise words so the distinctive product terms carry the Jaccard score."""
    text = (text or "").lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return {_stem(w) for w in text.split()} - _FUZZY_STOP_WORDS


def _fuzzy_score(a: str, b: str) -> float:
    wa, wb = normalize(a), normalize(b)
    if not wa or not wb:
        return 0.0
    return len(wa & wb) / len(wa | wb)


# ── URL mapping ──────────────────────────────────────────────────────────────


def _split_urls(cell) -> list[str]:
    """Split a (possibly comma-separated) cell into a list of clean, https:// URLs."""
    if not cell:
        return []
    text = str(cell).strip()
    if not text:
        return []
    urls = []
    for part in re.split(r"\s*,\s*", text):
        u = part.strip()
        if not u:
            continue
        if not re.match(r"^https?://", u, re.I):
            u = "https://" + u
        urls.append(u)
    return urls


def _load_url_mapping() -> dict:
    """Load data/url_mapping.xlsx into {SKU: {display_name, tenxyou, amazon, ajio, myntra}}.
    A SKU may span multiple rows (e.g. several TenXYou style variants sharing one
    SKU) — URLs from every matching row are pooled together into that SKU's lists."""
    mapping: dict = {}
    if not os.path.exists(URL_MAPPING_PATH):
        return mapping
    try:
        wb = openpyxl.load_workbook(URL_MAPPING_PATH, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = row[col["SKU"]]
            if not sku:
                continue
            sku = str(sku).strip()
            if not sku:
                continue
            entry = mapping.setdefault(sku, {
                "display_name": "",
                "tenxyou": [], "amazon": [], "ajio": [], "myntra": [],
            })
            display_name = row[col["TenXYou Display Name"]]
            if display_name and not entry["display_name"]:
                entry["display_name"] = str(display_name).strip()
            entry["tenxyou"].extend(_split_urls(row[col["TenXYou URL"]]))
            entry["amazon"].extend(_split_urls(row[col["AMAZON"]]))
            entry["ajio"].extend(_split_urls(row[col["AJIO"]]))
            entry["myntra"].extend(_split_urls(row[col["MYNTRA"]]))
    except Exception as e:
        print(f"[URL MAPPING ERROR] {URL_MAPPING_PATH}: {e}", flush=True)
        return {}
    return mapping


def _load_sku_mapping_fallback() -> dict:
    """Load data/sku_mapping.xlsx as a fallback source for SKUs absent from
    url_mapping.xlsx — its own Myntra/Ajio/TenXYou URL columns (there's no
    Amazon column on this older sheet) become that SKU's scrape URLs."""
    fallback: dict = {}
    if not os.path.exists(XLSX_PATH):
        return fallback
    try:
        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = row[col["SKU"]]
            if not sku:
                continue
            sku = str(sku).strip()
            fallback[sku] = {
                "display_name": str(row[col["Product Name"]] or "").strip(),
                "tenxyou": _split_urls(row[col["TenXYou URL"]]),
                "amazon":  [],
                "ajio":    _split_urls(row[col["Ajio URL"]]),
                "myntra":  _split_urls(row[col["Myntra URL"]]),
            }
    except Exception as e:
        print(f"[SKU MAPPING FALLBACK ERROR] {XLSX_PATH}: {e}", flush=True)
        return {}
    return fallback


# ── Core async scrape ──────────────────────────────────────────────────────────

MYNTRA_CACHE_PATH  = "data/myntra_search_results.json"
TENXYOU_CACHE_PATH = "data/tenxyou_search_results.json"
AJIO_CACHE_PATH    = "data/ajio_search_results.json"


def _load_cache(path: str) -> list:
    """Load products from a JSON cache file. Returns [] if missing or empty."""
    try:
        if not os.path.exists(path):
            return []
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) and data else []
    except Exception as e:
        print(f"[CACHE ERROR] {path}: {e}", flush=True)
        return []


def _build_mapped_url_set(url_mapping: dict, platform: str) -> set:
    """All URLs listed anywhere in url_mapping.xlsx for one platform, reduced to
    their stable IDs — keeps the fuzzy-match fallback from re-claiming a product
    that's already spoken for by an exact URL mapping (for this SKU or any
    other), robust to query params/color variants/slug wording differences."""
    return {_extract_stable_id(u, platform) for e in url_mapping.values() for u in e[platform]}


async def _scrape_one(sem: asyncio.Semaphore, browser, scrape_fn, sku: str, url: str) -> tuple:
    """Scrape a single (sku, url) pair under the shared concurrency semaphore.
    Returns (sku, url, price) — price is None on any failure."""
    async with sem:
        page = await browser.new_page()
        result = None
        try:
            result = await scrape_fn(page, url)
            price = _normalize_price(result.get("price")) if result.get("status") == "success" else None
        except Exception as e:
            print(f"[SCRAPE ERROR] {url}: {e}", flush=True)
            price = None
            result = {"status": f"exception: {e}"}
        finally:
            await page.close()

        if price is None:
            print(f"[SCRAPE FAIL] platform={scrape_fn.__name__} sku={sku} url={url} result={result}", flush=True)

        return sku, url, price


async def _scrape_all_lowest(sem: asyncio.Semaphore, browser, scrape_fn, sku_url_pairs: list) -> dict:
    """Scrape every (sku, url) pair concurrently (bounded by sem) and reduce to
    {sku: (lowest_price, matched_url)} — SKUs with no successful price are omitted.
    Used for Myntra/TenXYou individual-page scraping (asyncio.gather + semaphore)."""
    print(f"[SCRAPE_ALL_LOWEST] platform={scrape_fn.__name__} pairs={len(sku_url_pairs)}", flush=True)
    if not sku_url_pairs:
        return {}
    outcomes = await asyncio.gather(
        *(_scrape_one(sem, browser, scrape_fn, sku, url) for sku, url in sku_url_pairs)
    )
    best: dict = {}
    for sku, url, price in outcomes:
        if price is None:
            continue
        cur = best.get(sku)
        if cur is None or price < cur[0]:
            best[sku] = (price, url)
    return best


async def _scrape_amazon_lowest(context, urls: list) -> tuple:
    """Scrape every Amazon URL for one SKU sequentially and return
    (lowest_price, its_url). Returns (None, None) if urls is empty or none
    scraped a price successfully."""
    if not urls:
        return None, None
    best_price, best_url = None, None
    for url in urls:
        page = await context.new_page()
        try:
            result = await scrape_amazon_price_with_page(page, url)
            price = _normalize_price(result.get("price")) if result.get("status") == "success" else None
        except Exception as e:
            print(f"[SCRAPE ERROR] {url}: {e}", flush=True)
            price = None
        finally:
            await page.close()
        if price is not None and (best_price is None or price < best_price):
            best_price, best_url = price, url
    return best_price, best_url


def _lookup_ajio_lowest(urls: list, ajio_by_id: dict) -> tuple:
    """Ajio product pages are blocked by Akamai, so look prices up from the
    cached search-results JSON instead of scraping the page directly. Matches
    on the stable Ajio product code (ignoring color suffix and trailing '?') —
    ajio_by_id maps that code to *all* cached listings sharing it, since
    different color variants of the same product share a product code, and
    we still want the lowest price among them."""
    if not urls:
        return None, None
    best_price, best_url = None, None
    for url in urls:
        products = ajio_by_id.get(_extract_stable_id(url, "ajio"), [])
        for product in products:
            price = _normalize_price(product.get("price"))
            if price is not None and (best_price is None or price < best_price):
                best_price, best_url = price, url
    return best_price, best_url


def _lookup_tenxyou_exact(name: str, raw_tenxyou: list) -> tuple:
    """Cache fallback for TenXYou, mirroring how Ajio prices are looked up from
    cache: case-insensitive, whitespace-stripped exact name match against
    data/tenxyou_search_results.json. Returns (price, url), (None, None) if
    nothing matches."""
    target = (name or "").strip().lower()
    if not target:
        return None, None
    for p in raw_tenxyou:
        if (p.get("name") or "").strip().lower() == target:
            price = _normalize_price(p.get("price"))
            if price is not None:
                return price, p.get("url")
    return None, None


def _apply_supplementary_fuzzy_matching(results: list, raw_myntra: list, raw_tenxyou: list,
                                         excluded_urls: dict) -> None:
    """After primary per-URL scraping, fuzzy-match SKUs still missing a Myntra or
    TenXYou price against the cached search-results JSON (not a live scrape),
    skipping products already claimed by url_mapping.xlsx for any SKU. Mutates
    `results` in place, flagging a hit with "<platform>_fuzzy_match": True."""
    platform_specs = [
        ("myntra",  raw_myntra,  "myntra_price",  "matched_myntra_url",  "myntra_fuzzy_match"),
        ("tenxyou", raw_tenxyou, "tenxyou_price", "matched_tenxyou_url", "tenxyou_fuzzy_match"),
    ]

    for platform, raw_products, price_key, url_key, flag_key in platform_specs:
        excluded = excluded_urls.get(platform, set())
        candidates_pool = [p for p in raw_products if _extract_stable_id(p.get("url", ""), platform) not in excluded]
        if not candidates_pool:
            continue

        for r in results:
            if r.get(price_key) is not None:
                continue
            best_score, best_product = 0.0, None
            for p in candidates_pool:
                score = _fuzzy_score(r["name"], p.get("name", ""))
                if score > best_score:
                    best_score, best_product = score, p
            if best_product and best_score >= FUZZY_MATCH_THRESHOLD:
                price = _normalize_price(best_product.get("price"))
                if price is not None:
                    r[price_key] = price
                    r[url_key] = best_product["url"]
                    r[flag_key] = True


def _compute_status(r: dict) -> str:
    """complete — every platform with a mapped URL returned a price.
    partial — at least one price came back (URL-scraped or fuzzy-matched), but
    not every mapped platform succeeded.
    error — no prices at all.
    Platforms with no URL mapped (N/A) don't count against the SKU either way."""
    platforms = [
        (r["tenxyou_has_url"], r["tenxyou_price"]),
        (r["myntra_has_url"],  r["myntra_price"]),
        (r["ajio_has_url"],    r["ajio_price"]),
        (r["amazon_has_url"],  r["amazon_price"]),
    ]
    mapped = [price for has_url, price in platforms if has_url]
    got    = [price for price in mapped if price is not None]
    any_price = any(price is not None for _, price in platforms)

    if mapped and len(got) == len(mapped):
        return "complete"
    if any_price:
        return "partial"
    return "error"


async def _run_scrape(rows: list) -> list:
    """rows: [(sku, name), ...] from sku_mapping.xlsx — used as a display-name
    fallback and to pick up any SKUs that aren't in url_mapping.xlsx at all."""
    print("_RUN_SCRAPE CALLED", flush=True)

    url_mapping  = _load_url_mapping()
    sku_fallback = _load_sku_mapping_fallback()
    sheet_names  = {sku: name for sku, name in rows}

    raw_ajio = _load_cache(AJIO_CACHE_PATH)
    ajio_by_id: dict = defaultdict(list)
    for p in raw_ajio:
        ajio_by_id[_extract_stable_id(p.get("url", ""), "ajio")].append(p)

    print(f"[SCRAPE] URL mapping: {len(url_mapping)} SKUs loaded from {URL_MAPPING_PATH}", flush=True)
    print(f"[SCRAPE] Ajio cache:  {len(raw_ajio)} products loaded from {AJIO_CACHE_PATH}", flush=True)

    # url_mapping SKUs drive the loop; any sku_mapping-only SKU is appended as a fallback
    ordered_skus = list(url_mapping.keys())
    for sku in sheet_names:
        if sku not in url_mapping:
            ordered_skus.append(sku)

    entries = {
        sku: url_mapping.get(sku) or sku_fallback.get(sku) or {
            "display_name": "", "tenxyou": [], "amazon": [], "ajio": [], "myntra": [],
        }
        for sku in ordered_skus
    }

    # ── Primary: individual-page scraping. Myntra + TenXYou run concurrently
    # (bounded to SCRAPE_CONCURRENCY at a time); Amazon runs sequentially below.
    myntra_pairs  = [(sku, u) for sku in ordered_skus for u in entries[sku]["myntra"]]
    tenxyou_pairs = [(sku, u) for sku in ordered_skus for u in entries[sku]["tenxyou"]]
    sem = asyncio.Semaphore(SCRAPE_CONCURRENCY)

    async with async_playwright() as p:
        tenxyou_browser = await p.chromium.launch(headless=True)
        # headless=False: Myntra hangs/times out on headless Chrome navigations
        # (confirmed — headed mode succeeds in <1s, headless hangs to timeout).
        myntra_browser = await p.chromium.launch(headless=False, args=["--disable-http2"])
        amazon_browser = await p.chromium.launch(
            headless=False, args=["--disable-blink-features=AutomationControlled"]
        )
        amazon_context = await amazon_browser.new_context(
            user_agent=AMAZON_USER_AGENT, viewport={"width": 1920, "height": 1080}
        )

        try:
            myntra_best, tenxyou_best = await asyncio.gather(
                _scrape_all_lowest(sem, myntra_browser, scrape_myntra_price_with_page, myntra_pairs),
                _scrape_all_lowest(sem, tenxyou_browser, scrape_tenxyou_price_with_page, tenxyou_pairs),
            )

            results = []
            for sku in ordered_skus:
                entry = entries[sku]
                name = entry["display_name"] or sheet_names.get(sku, sku)

                tenxyou_price, matched_tenxyou_url = tenxyou_best.get(sku, (None, None))
                myntra_price,  matched_myntra_url  = myntra_best.get(sku, (None, None))
                ajio_price,    matched_ajio_url    = _lookup_ajio_lowest(entry["ajio"], ajio_by_id)
                amazon_price,  matched_amazon_url  = await _scrape_amazon_lowest(amazon_context, entry["amazon"])

                print(
                    f"  {sku:<12} T={tenxyou_price} M={myntra_price} A={ajio_price} AZ={amazon_price}",
                    flush=True,
                )

                results.append({
                    "sku":                 sku,
                    "name":                name,
                    "tenxyou_price":       tenxyou_price,
                    "myntra_price":        myntra_price,
                    "ajio_price":          ajio_price,
                    "amazon_price":        amazon_price,
                    "matched_tenxyou_url": matched_tenxyou_url,
                    "matched_myntra_url":  matched_myntra_url,
                    "matched_ajio_url":    matched_ajio_url,
                    "matched_amazon_url":  matched_amazon_url,
                    "tenxyou_has_url":     bool(entry["tenxyou"]),
                    "myntra_has_url":      bool(entry["myntra"]),
                    "ajio_has_url":        bool(entry["ajio"]),
                    "amazon_has_url":      bool(entry["amazon"]),
                    "tenxyou_fuzzy_match": False,
                    "myntra_fuzzy_match":  False,
                    "ajio_fuzzy_match":    False,
                    "amazon_fuzzy_match":  False,
                })
        finally:
            await tenxyou_browser.close()
            await myntra_browser.close()
            await amazon_browser.close()

    # ── Supplementary: fuzzy-match still-missing Myntra/TenXYou prices against
    # cached (not live) search results.
    print("[SCRAPE] Loading cached search results for supplementary fuzzy matching...", flush=True)
    raw_myntra_cache  = _load_cache(MYNTRA_CACHE_PATH)
    raw_tenxyou_cache = _load_cache(TENXYOU_CACHE_PATH)

    # TenXYou cache fallback: exact name match first (cheap, high-confidence,
    # same idea as the Ajio cache lookup) before the broader fuzzy word-overlap pass.
    for r in results:
        if r["tenxyou_price"] is None:
            price, url = _lookup_tenxyou_exact(r["name"], raw_tenxyou_cache)
            if price is not None:
                r["tenxyou_price"] = price
                r["matched_tenxyou_url"] = url
                r["tenxyou_fuzzy_match"] = True

    excluded_urls = {
        "myntra":  _build_mapped_url_set(url_mapping, "myntra"),
        "tenxyou": _build_mapped_url_set(url_mapping, "tenxyou"),
    }
    _apply_supplementary_fuzzy_matching(results, raw_myntra_cache, raw_tenxyou_cache, excluded_urls)

    for r in results:
        comp_prices = [p for p in [r["myntra_price"], r["ajio_price"], r["amazon_price"]] if p is not None]
        r["lowest_comp_price"] = min(comp_prices) if comp_prices else None
        r["status"] = _compute_status(r)

    return results


# ── Flask routes ───────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/scrape", methods=["POST"])
def scrape():
    print("SCRAPE ROUTE CALLED", flush=True)
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    sku_col  = headers.index("SKU")
    name_col = headers.index("Product Name")

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[sku_col]
        if not sku:
            continue
        rows.append((sku, (row[name_col] or "").strip()))

    results = asyncio.run(_run_scrape(rows))

    output = {
        "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "count": len(results),
        "results": results,
    }
    with open(RESULTS_PATH, "w") as f:
        json.dump(output, f, indent=2)

    return jsonify(output)


@app.route("/refresh-ajio", methods=["POST"])
def refresh_ajio():
    products = _load_cache(AJIO_CACHE_PATH)
    return jsonify({
        "count": len(products),
        "note": "Run `python scrapers/ajio_search.py` from terminal to refresh Ajio data.",
    })


@app.route("/results")
def results():
    if not os.path.exists(RESULTS_PATH):
        return jsonify({"timestamp": None, "count": 0, "results": []})
    with open(RESULTS_PATH) as f:
        return jsonify(json.load(f))


@app.route("/export")
def export():
    if not os.path.exists(RESULTS_PATH):
        return make_response("No results to export", 404)
    with open(RESULTS_PATH) as f:
        data = json.load(f)

    si = StringIO()
    writer = csv.DictWriter(
        si,
        fieldnames=["sku", "name", "tenxyou_price", "myntra_price", "ajio_price", "amazon_price",
                    "lowest_comp_price",
                    "matched_tenxyou_url", "matched_myntra_url", "matched_ajio_url", "matched_amazon_url",
                    "tenxyou_has_url", "myntra_has_url", "ajio_has_url", "amazon_has_url",
                    "tenxyou_fuzzy_match", "myntra_fuzzy_match", "ajio_fuzzy_match", "amazon_fuzzy_match",
                    "status"],
        extrasaction="ignore",
    )
    writer.writeheader()
    writer.writerows(data["results"])

    response = make_response(si.getvalue())
    response.headers["Content-Disposition"] = "attachment; filename=price_results.csv"
    response.headers["Content-Type"] = "text/csv"
    return response


if __name__ == "__main__":
    app.run(debug=True, use_reloader=False)
