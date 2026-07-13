import asyncio
import json
import logging
import os
import re
import sys
from urllib.parse import parse_qs, unquote, urlparse

import openpyxl
from playwright.async_api import async_playwright

# Ensure project root is on sys.path so imports/paths work from any CWD
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

STORE_URL       = "https://www.amazon.in/stores/TenxYou/page/A64FEC8A-4DF3-4B89-90C0-4FA6B50C5F5F"
SEARCH_URL      = "https://www.amazon.in/s?k=ten+x+you"
RESULTS_PATH    = os.path.join(_ROOT, "data", "amazon_search_results.json")
SCREENSHOT_PATH = os.path.join(_ROOT, "data", "amazon_debug_screenshot.png")
SEARCH_SCREENSHOT_PATH = os.path.join(_ROOT, "data", "amazon_debug_screenshot_search.png")
LOG_PATH        = os.path.join(_ROOT, "data", "debug_scrape.log")
URL_MAPPING_PATH = os.path.join(_ROOT, "data", "url_mapping.xlsx")

SEARCH_ITEM_SELECTOR = 'div[data-component-type="s-search-result"]'

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

CAPTCHA_MARKERS = [
    "enter the characters you see below",
    "type the characters you see in this image",
    "api-services-support@amazon.com",
    "sorry, we just need to make sure you're not a robot",
    "to discuss automated access to amazon data",
    "robot check",
]

logger = logging.getLogger("amazon_search")
logger.setLevel(logging.INFO)
logger.propagate = False
if not logger.handlers:
    _file_handler = logging.FileHandler(LOG_PATH, encoding="utf-8")
    _file_handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    _stream_handler = logging.StreamHandler()
    _stream_handler.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(_file_handler)
    logger.addHandler(_stream_handler)


async def _extract_products(page) -> list:
    """Pull product tiles from an Amazon store page via /dp/ product links,
    since store-page layouts vary but product links are always /dp/<ASIN>."""
    products = []
    seen = set()
    anchors = await page.query_selector_all("a[href*='/dp/']")

    for a in anchors:
        href = await a.get_attribute("href")
        if not href:
            continue
        url = href if href.startswith("http") else f"https://www.amazon.in{href}"
        url = url.split("?")[0]

        m = re.search(r"/dp/([A-Z0-9]{10})", url)
        key = m.group(1) if m else url
        if key in seen:
            continue
        seen.add(key)

        name = await a.get_attribute("aria-label") or ""
        if not name:
            img = await a.query_selector("img")
            if img:
                name = await img.get_attribute("alt") or ""
        if not name:
            name = (await a.inner_text()).strip()
        name = name.strip() or "Not found"

        price = "Not found"
        node = a
        for _ in range(4):
            handle = await node.evaluate_handle("el => el.parentElement")
            parent = handle.as_element()
            if not parent:
                break
            node = parent
            text = await node.inner_text()
            match = re.search(r"₹\s?[\d,]+(?:\.\d+)?", text)
            if match:
                price = match.group(0)
                break

        products.append({"name": name, "price": price, "url": url})

    return products


def _resolve_amazon_url(href: str) -> str:
    """Sponsored results link through /sspa/click?...&url=<encoded real path>;
    unwrap that redirect to get the actual product URL."""
    if href.startswith("/sspa/click"):
        real = parse_qs(urlparse(href).query).get("url", [None])[0]
        if real:
            href = unquote(real)
    url = href if href.startswith("http") else f"https://www.amazon.in{href}"
    return url.split("?")[0]


async def _extract_search_results(page) -> list:
    """Pull product cards from an Amazon search results grid. Each card's
    title lives in an <a><h2 aria-label="full title">...</h2></a> inside
    [data-cy="title-recipe"] — a sibling brand-only <h2> comes before it."""
    products = []
    seen = set()
    items = await page.query_selector_all(SEARCH_ITEM_SELECTOR)

    for item in items:
        anchor = await item.query_selector('[data-cy="title-recipe"] a')
        href = await anchor.get_attribute("href") if anchor else None
        if not href or href.startswith("javascript:"):
            continue
        url = _resolve_amazon_url(href)

        m = re.search(r"/dp/([A-Z0-9]{10})", url)
        key = m.group(1) if m else url
        if key in seen:
            continue
        seen.add(key)

        title_h2 = await anchor.query_selector("h2")
        name = await title_h2.get_attribute("aria-label") if title_h2 else None
        if not name and title_h2:
            name = (await title_h2.inner_text()).strip()
        name = (name or "Not found").strip()

        price_el = await item.query_selector(".a-price .a-offscreen")
        price = (await price_el.inner_text()).strip() if price_el else "Not found"

        products.append({"name": name, "price": price, "url": url})

    return products


def _asin(url: str) -> str:
    m = re.search(r"/dp/([A-Z0-9]{10})", url)
    return m.group(1) if m else url


def _load_known_amazon_asins() -> set:
    """Load confirmed TenXYou ASINs from data/url_mapping.xlsx's AMAZON column —
    the ground-truth source of known-good Amazon listings for this brand."""
    asins = set()
    if not os.path.exists(URL_MAPPING_PATH):
        return asins
    try:
        wb = openpyxl.load_workbook(URL_MAPPING_PATH, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        if "AMAZON" not in headers:
            return asins
        col = headers.index("AMAZON")
        for row in ws.iter_rows(min_row=2, values_only=True):
            cell = row[col]
            if not cell:
                continue
            for part in str(cell).split(","):
                m = re.search(r"/dp/([A-Z0-9]{10})", part)
                if m:
                    asins.add(m.group(1))
    except Exception as e:
        logger.error(f"Error loading known ASINs from {URL_MAPPING_PATH}: {e}")
    return asins


LOOSE_RELEVANCE_TERMS = [
    "ten x you", "sachin tendulkar", "xu0", "xm1", "xw1", "xa1",
    "unisex reset", "unisex pivot", "unisex crossover", "unisex aeonic",
    "unisex zenflo", "unisex sundowner", "unisex switch", "unisex nox",
    "centurion", "youngstar",
]


def _matches_loose_terms(name: str) -> bool:
    lowered = (name or "").lower()
    return any(term in lowered for term in LOOSE_RELEVANCE_TERMS)


def _is_known_tenxyou_product(prod: dict, known_asins: set) -> bool:
    """A product counts as a genuine TenXYou listing if its ASIN is already
    confirmed in url_mapping.xlsx, or — for ASINs not yet mapped — if its name
    matches one of the known TenXYou product-line/brand terms."""
    if _asin(prod.get("url", "")) in known_asins:
        return True
    return _matches_loose_terms(prod.get("name", ""))


async def scrape_amazon_search_page(search_url: str = SEARCH_URL) -> list:
    """Scrape exactly 2 pages of Amazon search results (?page=1, ?page=2),
    combine, and deduplicate by ASIN."""
    all_products = []
    page_counts = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )
        context = await browser.new_context(
            user_agent=USER_AGENT,
            viewport={"width": 1920, "height": 1080},
        )
        page = await context.new_page()

        try:
            for page_num in (1, 2):
                paged_url = f"{search_url}&page={page_num}"
                logger.info(f"Navigating to {paged_url}")
                await page.goto(paged_url, wait_until="domcontentloaded", timeout=30000)

                try:
                    await page.wait_for_load_state("networkidle", timeout=20000)
                except Exception:
                    logger.info("Network did not go fully idle within timeout, continuing anyway")

                await page.wait_for_timeout(2000)

                screenshot_path = SEARCH_SCREENSHOT_PATH.replace(".png", f"_page{page_num}.png")
                os.makedirs(os.path.dirname(screenshot_path), exist_ok=True)
                await page.screenshot(path=screenshot_path, full_page=True)
                logger.info(f"Saved debug screenshot to {screenshot_path}")

                content = (await page.content()).lower()
                if any(marker in content for marker in CAPTCHA_MARKERS):
                    logger.error(
                        f"CAPTCHA / bot-check page detected on Amazon search page {page_num} - not retrying. "
                        f"See {screenshot_path} to inspect what rendered."
                    )
                    page_counts.append((page_num, 0))
                    continue

                page_products = await _extract_search_results(page)
                page_counts.append((page_num, len(page_products)))
                logger.info(f"Page {page_num}: {len(page_products)} products")

                for prod in page_products:
                    logger.info(f"{prod['name']} | {prod['price']} | {prod['url']}")

                all_products.extend(page_products)

        except Exception as e:
            logger.error(f"Error scraping Amazon search results: {e}")

        finally:
            await browser.close()

    seen = set()
    products = []
    for prod in all_products:
        key = _asin(prod["url"])
        if key not in seen:
            seen.add(key)
            products.append(prod)

    known_asins = _load_known_amazon_asins()
    logger.info(f"Loaded {len(known_asins)} known TenXYou ASINs from {URL_MAPPING_PATH}")

    before_relevance_filter = len(products)
    products = [p for p in products if _is_known_tenxyou_product(p, known_asins)]
    logger.info(f"Relevance filter: kept {len(products)} of {before_relevance_filter}")

    before_price_filter = len(products)
    products = [p for p in products if p.get("price") != "Not found"]
    no_price_removed = before_price_filter - len(products)

    with open(RESULTS_PATH, "w", encoding="utf-8") as f:
        json.dump(products, f, indent=2, ensure_ascii=False)

    for page_num, count in page_counts:
        logger.info(f"Page {page_num}: {count} products (raw)")
    logger.info(f"Removed {no_price_removed} products with no price found")
    logger.info(f"Total unique: {len(products)} products. Saved to {RESULTS_PATH}")
    logger.info("First 5 products:")
    for prod in products[:5]:
        logger.info(f"  {prod['name']}")

    return products


async def scrape_amazon_search(store_url: str = STORE_URL) -> list:
    products = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )
        context = await browser.new_context(
            user_agent=USER_AGENT,
            viewport={"width": 1920, "height": 1080},
        )
        page = await context.new_page()

        try:
            logger.info(f"Navigating to {store_url}")
            await page.goto(store_url, wait_until="domcontentloaded", timeout=30000)

            try:
                await page.wait_for_load_state("networkidle", timeout=20000)
            except Exception:
                logger.info("Network did not go fully idle within timeout, continuing anyway")

            await page.wait_for_timeout(2000)

            os.makedirs(os.path.dirname(SCREENSHOT_PATH), exist_ok=True)
            await page.screenshot(path=SCREENSHOT_PATH, full_page=True)
            logger.info(f"Saved debug screenshot to {SCREENSHOT_PATH}")

            content = (await page.content()).lower()
            is_captcha = any(marker in content for marker in CAPTCHA_MARKERS)

            if is_captcha:
                logger.error(
                    "CAPTCHA / bot-check page detected on Amazon store page - not retrying. "
                    f"See {SCREENSHOT_PATH} to inspect what rendered."
                )
                return []

            products = await _extract_products(page)

            if not products:
                logger.error(
                    "No product content found on the Amazon store page (blocked or empty) - not retrying. "
                    f"See {SCREENSHOT_PATH} to inspect what rendered."
                )
                return []

            for prod in products:
                logger.info(f"{prod['name']} | {prod['price']} | {prod['url']}")

        except Exception as e:
            logger.error(f"Error scraping Amazon store page: {e}")
            return []

        finally:
            await browser.close()

    with open(RESULTS_PATH, "w", encoding="utf-8") as f:
        json.dump(products, f, indent=2, ensure_ascii=False)

    logger.info(f"Scraped {len(products)} products. Saved to {RESULTS_PATH}")
    return products


async def scrape_amazon_price_with_page(page, url: str) -> dict:
    """Scrape name + price off an individual Amazon product (dp) page.
    Same {url, name, price, status} shape as scrapers/tenxyou.py and
    scrapers/myntra.py's single-product scrapers, so app.py can treat all
    three platforms uniformly."""
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=30000)
        await page.wait_for_timeout(2000)

        content = (await page.content()).lower()
        if any(marker in content for marker in CAPTCHA_MARKERS):
            return {"url": url, "name": None, "price": None, "status": "error: captcha"}

        # Scoped to Amazon's dedicated #availability element rather than the whole
        # page: "currently unavailable" also shows up in generic JS translation
        # strings and other-swatch status text on virtually every multi-variant
        # listing, which false-positives a page-wide substring search even for
        # products that are genuinely in stock and correctly priced.
        availability_el = await page.query_selector("#availability")
        availability_text = (await availability_el.inner_text()).strip().lower() if availability_el else ""
        if "currently unavailable" in availability_text or "temporarily out of stock" in availability_text:
            return {"url": url, "name": None, "price": None, "status": "unavailable"}

        name_el = await page.query_selector("#productTitle")
        name = (await name_el.inner_text()).strip() if name_el else "Not found"

        price = None
        for sel in (
            "#corePrice_feature_div .a-price .a-offscreen",
            "#corePriceDisplay_desktop_feature_div span.a-price-whole",
            "#priceblock_ourprice",
            "#priceblock_dealprice",
        ):
            el = await page.query_selector(sel)
            if el:
                text = (await el.inner_text()).strip()
                if text:
                    price = text
                    break
        if price is None:
            for el in await page.query_selector_all(".a-price .a-offscreen"):
                text = (await el.inner_text()).strip()
                if text:
                    price = text
                    break
        price = price or "Not found"

        return {"url": url, "name": name, "price": price, "status": "success"}

    except Exception as e:
        return {"url": url, "name": None, "price": None, "status": f"error: {str(e)}"}


async def scrape_amazon_price(url: str) -> dict:
    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )
        context = await browser.new_context(
            user_agent=USER_AGENT,
            viewport={"width": 1920, "height": 1080},
        )
        page = await context.new_page()
        try:
            return await scrape_amazon_price_with_page(page, url)
        finally:
            await browser.close()


if __name__ == "__main__":
    asyncio.run(scrape_amazon_search_page())
