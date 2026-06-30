import asyncio
import openpyxl
from scrapers.myntra import scrape_myntra_price


async def main():
    wb = openpyxl.load_workbook("data/sku_mapping.xlsx")
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    myntra_col = headers.index("Myntra URL")

    for row in ws.iter_rows(min_row=2, values_only=True):
        myntra_url = row[myntra_col]
        if not myntra_url:
            continue

        sku = row[headers.index("SKU")]
        print(f"\nScraping SKU: {sku} — {myntra_url}")
        result = await scrape_myntra_price(myntra_url)
        print(result)


if __name__ == "__main__":
    asyncio.run(main())
